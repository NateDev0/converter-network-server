"""images_to_pptx.py <dir-of-page-pngs> <out.pptx>: one full-bleed slide per page."""
import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

src, out = Path(sys.argv[1]), sys.argv[2]
pages = sorted((int(m.group(1)), p) for p in src.iterdir() if (m := re.search(r"(\d+)\.png$", p.name)))
if not pages:
    sys.exit("no pages")
w, h = Image.open(pages[0][1]).size
prs = Presentation()
prs.slide_width = Emu(12192000)  # 13.333 in (16:9 width)
prs.slide_height = Emu(int(12192000 * h / w))
blank = prs.slide_layouts[6]
for _, p in pages:
    prs.slides.add_slide(blank).shapes.add_picture(str(p), 0, 0, prs.slide_width, prs.slide_height)
prs.save(out)
